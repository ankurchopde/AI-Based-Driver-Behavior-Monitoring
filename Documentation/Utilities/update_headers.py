import docx
from docx.shared import Inches

def update_headers(output_path):
    doc = docx.Document(output_path)
    
    doc.paragraphs[21].text = "Current Progress Summary"
    doc.paragraphs[23].text = "Problems and Solutions"
    doc.paragraphs[25].text = "Simulation Results Screenshot"
    doc.paragraphs[28].text = "Next Week's Work Plan"
    doc.paragraphs[30].text = "Expected Next-Level Model"
    doc.paragraphs[32].text = "Observations and Recommendations"
    doc.paragraphs[34].text = "Conclusion of This Week"
    
    doc.save(output_path)

if __name__ == '__main__':
    update_headers('WEEK_3_PROJECT_PROGRESS_REPORT_AE.docx')
