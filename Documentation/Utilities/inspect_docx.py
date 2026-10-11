import docx

def read_template(filepath):
    doc = docx.Document(filepath)
    print("--- PARAGRAPHS ---")
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            print(f"Para {i}: {p.text}")
    print("\n--- TABLES ---")
    for t_idx, table in enumerate(doc.tables):
        print(f"Table {t_idx}:")
        for r_idx, row in enumerate(table.rows):
            row_data = [cell.text.strip() for cell in row.cells]
            print(f"  Row {r_idx}: {row_data}")

if __name__ == '__main__':
    read_template('WEEKLY_PROJECT_PROGRESS_REPORT_TEMPLET_AE.docx')
