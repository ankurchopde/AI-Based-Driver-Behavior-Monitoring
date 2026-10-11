import docx

def fix_last_details(output_path):
    doc = docx.Document(output_path)
    
    doc.paragraphs[19].text = "Simulation / Testing"
    
    # fix the degree symbol issue
    if "30" in doc.paragraphs[16].text:
        doc.paragraphs[16].text = doc.paragraphs[16].text.replace("30°", "30 degrees").replace("20°", "20 degrees")
        
    doc.save(output_path)

if __name__ == '__main__':
    fix_last_details('WEEK_3_PROJECT_PROGRESS_REPORT_AE.docx')
