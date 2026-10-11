import docx
from docx.shared import Inches

def create_report(template_path, output_path, image_path):
    doc = docx.Document(template_path)
    
    # 1. Fill Table 0 (Metadata)
    t0 = doc.tables[0]
    t0.rows[0].cells[1].text = "AI-Based Driver Behavior Monitoring System"
    t0.rows[2].cells[1].text = "Week 3"
    t0.rows[3].cells[1].text = "Tanmaygt1 (Lead)"
    t0.rows[4].cells[1].text = "Project Supervisor"
    t0.rows[5].cells[1].text = "Automotive / Software Engineering"
    t0.rows[6].cells[1].text = "August 2026"
    
    # 2. Fill Paragraphs
    doc.paragraphs[2].text = "AI-Based Driver Behavior Monitoring System (Software-In-the-Loop Prototype)"
    
    doc.paragraphs[5].text = "The primary objective of Week 3 was the completion and functional simulation of the integrated MATLAB/Simulink driver monitoring prototype. This involved combining the previously developed Drowsiness, Distraction, and Driver Identification components into a single Simulink environment, resolving integration challenges, and conducting functional testing."
    
    doc.paragraphs[9].text = (
        "1. Integrated Section 1 (Drowsiness), Section 2 (Distraction), and Section 3 (Driver Identification) into a final MATLAB/Simulink dashboard.\n"
        "2. Configured custom UDP telemetry (88-byte packet) to transmit 11 spatial/state variables at exactly 20 Hz from Python to MATLAB.\n"
        "3. Configured Stateflow/Simulink temporal persistence filters (e.g., 2-second delay for drowsiness) to avoid false-positive alerts.\n"
        "4. Conducted functional testing using open-source datasets and local webcam feeds."
    )
    
    doc.paragraphs[14].text = (
        "The final architecture consists of a Python Vision Node that captures webcam frames and computes metrics using MediaPipe and ONNX SFace in <50ms. A custom MATLAB System Object (UDPReceiverSystemIdentification.m) unpacks the byte stream into 11 distinct signal lines in Simulink, which are then passed through temporal logic gates to trigger actionable alerts."
    )
    
    doc.paragraphs[15].text = "Drowsiness, Distraction & Identification Implementation"
    doc.paragraphs[16].text = (
        "- Drowsiness Detection: Tracks Eye Aspect Ratio (EAR < 0.20) and Mouth Aspect Ratio (MAR > 0.60) with a 2-second persistence filter to ignore normal blinks.\n"
        "- Distraction Detection: Analyzes Head Pose (Yaw > 30°, Pitch > 20°) using OpenCV solvePnP on the 3D face mesh, also gated by a 2-second persistence timer.\n"
        "- Driver Identification: Utilizes an ONNX face recognizer to compare real-time embeddings against a local database, using a 0.363 cosine similarity threshold and 10-frame persistence."
    )
    
    doc.paragraphs[17].text = "Alert / Decision Logic"
    doc.paragraphs[18].text = (
        "The logic combines raw signals with state transitions. The system only registers an unsafe state if the condition persists beyond the predefined temporal threshold, at which point an alert is output on the Simulink dashboard."
    )
    
    # Modify Table 1 (Test Cases)
    t1 = doc.tables[1]
    # Update headers
    t1.rows[0].cells[0].text = "Test Scenario"
    t1.rows[0].cells[1].text = "Detected State"
    t1.rows[0].cells[2].text = "Expected Output"
    t1.rows[0].cells[3].text = "Observed Output"
    t1.rows[0].cells[4].text = "Status"
    
    # We need to add rows
    tests = [
        ["Prolonged Eye Closure", "EAR < 0.20 for 2s", "Drowsiness Alert", "Triggers at 2.0s mark", "PASS"],
        ["Head turned right", "Yaw > 30 deg for 2s", "Distraction Alert", "Triggers on lateral move", "PASS"],
        ["Normal Blink", "EAR < 0.20 for 0.5s", "No Alert", "Filtered by logic", "PASS"],
        ["Known Driver", "Similarity > 0.363", "Access Granted", "Validated ID", "PASS"]
    ]
    # Fill row 1
    for i, val in enumerate(tests[0]):
        t1.rows[1].cells[i].text = val
    # Add remaining rows
    for test in tests[1:]:
        row_cells = t1.add_row().cells
        for i, val in enumerate(test):
            row_cells[i].text = val
            
    # Modify Table 2 (Progress)
    t2 = doc.tables[2]
    progress = [
        ["Level 1", "Drowsiness Module", "Completed"],
        ["Level 2", "Distraction Module", "Completed"],
        ["Level 3", "Driver Identification", "Completed"],
        ["Level 4", "UDP Integration & Simulink Dashboard", "Completed"],
        ["Level 5", "Edge-AI Hardware Deployment (Raspberry Pi)", "Not Started"],
        ["Level 6", "Real Vehicle Hardware Testing", "Not Started"]
    ]
    for r_idx in range(1, 7):
        for c_idx in range(3):
            t2.rows[r_idx].cells[c_idx].text = progress[r_idx-1][c_idx]
            
    # Problems and Solutions
    doc.paragraphs[24].text = (
        "Problem: Simulink System Object paths failing to resolve on startup.\n"
        "Cause: Missing bindings for the custom MATLAB System Object cross-section dependencies.\n"
        "Solution: A startup.m script was implemented to ensure paths are properly bound before opening the Simulink model.\n"
        "Status: Resolved."
    )
    
    # Image replacement
    doc.paragraphs[26].text = "" # Clear placeholder text
    run = doc.paragraphs[26].add_run()
    run.add_picture(image_path, width=Inches(6.0))
    
    doc.paragraphs[27].text = "Figure 1: Integrated MATLAB/Simulink Driver Monitoring Model Dashboard and 11-value UDP telemetry scopes."
    
    # Modify Table 3 (Next Week)
    t3 = doc.tables[3]
    next_tasks = [
        ["Deploy system to Edge-AI hardware (Raspberry Pi)", "High"],
        ["Optimize deep learning models for low-compute environments", "High"],
        ["Integrate with physical CAN-bus hardware alerts", "Medium"]
    ]
    for r_idx, task in enumerate(next_tasks, start=1):
        t3.rows[r_idx].cells[1].text = task[0]
        t3.rows[r_idx].cells[2].text = task[1]
    # Clear unused rows
    for r_idx in range(4, 9):
        t3.rows[r_idx].cells[1].text = ""
        t3.rows[r_idx].cells[2].text = ""
        
    doc.paragraphs[31].text = "The expected next-level model involves an Edge-AI accelerated device (Raspberry Pi 4/5) with TFLite/ONNX deployment running independent of the PC, signaling physical hardware (e.g., haptic seat vibrations) instead of Software-In-the-Loop scopes."
    
    doc.paragraphs[33].text = "While Software-In-the-Loop functional simulation successfully validates the deterministic logic, real-world conditions introduce environmental variance. The SFace cosine similarity threshold (0.363) and illumination tolerance must be calibrated under actual automotive cabin lighting conditions."
    
    doc.paragraphs[35].text = "Week 3 successfully delivered the integrated simulation prototype. All three main vision subsystems communicate reliably with Simulink at 20 Hz, satisfying the project's software integration and functional simulation objectives."
    
    doc.save(output_path)

if __name__ == '__main__':
    template = 'WEEKLY_PROJECT_PROGRESS_REPORT_TEMPLET_AE.docx'
    output = 'WEEK_3_PROJECT_PROGRESS_REPORT_AE.docx'
    image = r'd:\Academics\TY_SEM1\AE\AI based Driver Detector\AI-Based-Driver-Behavior-Monitoring\Results\Final_Integration\01_udp_11_values.png'
    create_report(template, output, image)
    print(f"Successfully saved to {output}")
